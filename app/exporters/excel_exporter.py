"""
Excel exporter for the lead generation pipeline.

Creates a professionally formatted .xlsx file with:
  - Styled header row with filters
  - Frozen top row
  - Auto-adjusted column widths
  - Color-coded headers by data group
"""

from pathlib import Path
from typing import List

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from app.core.models import BusinessLead
from app.core.logger import setup_logger


logger = setup_logger("excel_exporter")

# Column order and header display names
COLUMN_CONFIG = [
    # (field_name, display_header, group_color)
    ("business_name", "Business Name", "4472C4"),
    ("business_category", "Category", "4472C4"),
    ("industry", "Industry", "4472C4"),
    ("address", "Full Address", "4472C4"),
    ("street", "Street", "4472C4"),
    ("city", "City", "4472C4"),
    ("state", "State", "4472C4"),
    ("zip_code", "ZIP Code", "4472C4"),
    ("phone", "Phone (Maps)", "548235"),
    ("phone_from_website", "Phone (Website)", "548235"),
    ("website", "Website", "548235"),
    ("email", "Email", "548235"),
    ("google_rating", "Rating", "BF8F00"),
    ("review_count", "Reviews", "BF8F00"),
    ("business_status", "Status", "BF8F00"),
    ("price_level", "Price Level", "BF8F00"),
    ("opening_hours", "Hours", "BF8F00"),
    ("facebook_url", "Facebook", "7030A0"),
    ("instagram_url", "Instagram", "7030A0"),
    ("linkedin_url", "LinkedIn", "7030A0"),
    ("youtube_url", "YouTube", "7030A0"),
    ("has_social_presence", "Has Socials", "7030A0"),
    ("contact_page", "Contact Page", "7030A0"),
    ("booking_url", "Booking URL", "7030A0"),
    ("google_maps_url", "Maps URL", "808080"),
    ("latitude", "Latitude", "808080"),
    ("longitude", "Longitude", "808080"),
    ("extraction_status", "Status", "808080"),
    ("data_collected_at", "Collected At", "808080"),
]


def export_excel(leads: List[BusinessLead], output_path: str) -> str:
    """
    Export leads to a professionally formatted Excel file.

    Args:
        leads: List of cleaned/validated BusinessLead records
        output_path: Path for the output .xlsx file

    Returns:
        Path to the created file
    """
    Path(output_path).parent.mkdir(parents=True, exist_ok=True)

    # Build DataFrame
    field_names = [c[0] for c in COLUMN_CONFIG]
    display_names = [c[1] for c in COLUMN_CONFIG]

    rows = []
    for lead in leads:
        data = lead.model_dump()
        rows.append({c[1]: data.get(c[0]) for c in COLUMN_CONFIG})

    df = pd.DataFrame(rows, columns=display_names)
    df.to_excel(output_path, index=False, sheet_name="Leads", engine="openpyxl")

    # ── Apply formatting ────────────────────────────────────────────────
    wb = load_workbook(output_path)
    ws = wb["Leads"]

    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    # Style headers
    for col_idx, (field, display, color) in enumerate(COLUMN_CONFIG, 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = Font(bold=True, color="FFFFFF", size=10)
        cell.fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    # Style data rows
    for row_idx in range(2, ws.max_row + 1):
        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.alignment = Alignment(vertical="top", wrap_text=False)
            cell.border = thin_border

    # Auto-width columns (with min/max constraints)
    for col_idx in range(1, ws.max_column + 1):
        max_length = 0
        col_letter = get_column_letter(col_idx)
        for row in ws.iter_rows(min_col=col_idx, max_col=col_idx, values_only=True):
            for cell_value in row:
                if cell_value:
                    max_length = max(max_length, len(str(cell_value)))
        adjusted_width = min(max(max_length + 2, 10), 40)
        ws.column_dimensions[col_letter].width = adjusted_width

    # Freeze top row
    ws.freeze_panes = "A2"

    # Add auto-filter
    ws.auto_filter.ref = ws.dimensions

    wb.save(output_path)

    logger.info(f"Excel exported: {len(leads)} records -> {output_path}")
    return output_path
